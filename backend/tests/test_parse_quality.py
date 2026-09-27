"""Increment A — parsing robustness and the parse-status guard.

Each test here corresponds to a bug that was reproduced against the running
code before the fix (see AI-QUALITY-PLAN.md §1).
"""

import io

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app.services.exceptions import ResumeNotReady
from app.services.parsing import MIN_WORDS_PDF, extract_text, structure_text


# ----------------------------------------------------------- §3.6 numbered
def test_numbered_bullets_are_recognised():
    """Before: `1.` / `2)` lists produced ZERO experience entries."""
    data = structure_text(
        "EXPERIENCE\n"
        "Acme Corp - Senior Engineer - 2020-2024\n"
        "1. Led the platform migration to Kubernetes.\n"
        "2) Built Terraform modules adopted by 8 teams.\n"
    )
    assert len(data["experience"]) == 1
    entry = data["experience"][0]
    assert entry["company"] == "Acme Corp"
    assert entry["bullets"] == [
        "Led the platform migration to Kubernetes.",
        "Built Terraform modules adopted by 8 teams.",
    ]


def test_glyph_bullets_still_work():
    """The numbered-bullet change must not regress the original formats."""
    data = structure_text(
        "EXPERIENCE\n"
        "Acme Corp - Senior Engineer - 2020-2024\n"
        "- Led the migration.\n"
        "\u2022 Built modules.\n"
        "* Shipped the thing.\n"
    )
    assert data["experience"][0]["bullets"] == [
        "Led the migration.",
        "Built modules.",
        "Shipped the thing.",
    ]


@pytest.mark.parametrize(
    "line",
    [
        "3.5x faster than the previous pipeline",  # decimal, no space after the dot
        "2024 - Present",                          # a bare date line
        "1000 requests per second sustained",      # a leading metric
    ],
)
def test_numeric_prose_is_not_mistaken_for_a_list_marker(line):
    """The numbered-bullet regex requires `<digits><.|)><space>`.

    Asserted against _BULLET_RE directly: routed through structure_text these
    lines hit the glyph branch or the header branch first, so the test would
    pass without ever exercising the numeric alternative.
    """
    from app.services.parsing import _BULLET_RE

    assert _BULLET_RE.match(line) is None


@pytest.mark.parametrize("line", ["1. Did a thing", "2) Did another", "10.  Spaced out"])
def test_real_numbered_markers_do_match(line):
    from app.services.parsing import _BULLET_RE

    assert _BULLET_RE.match(line) is not None


# ------------------------------------------------------------- §3.4 tables
def _docx_with_table() -> bytes:
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("SKILLS")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Python, Go, Rust"
    table.cell(1, 0).text = "Cloud"
    table.cell(1, 1).text = "AWS, Kubernetes, Terraform"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.parametrize("skill", ["Python", "Go", "Rust", "AWS", "Kubernetes", "Terraform"])
def test_docx_table_content_survives_extraction(skill):
    """Before: document.paragraphs excluded table cells, losing the whole grid."""
    assert skill in extract_text("resume.docx", _docx_with_table())


def test_docx_paragraphs_still_extracted():
    assert "Jane Doe" in extract_text("resume.docx", _docx_with_table())


# --------------------------------------------------------- §3.2 scanned PDF
def _image_only_pdf() -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(100, 500, 300, 200, fill=1)  # no text operators at all
    c.showPage()
    c.save()
    return buf.getvalue()


def test_image_only_pdf_extracts_no_text():
    assert len(extract_text("scan.pdf", _image_only_pdf()).split()) < MIN_WORDS_PDF


def test_scanned_pdf_is_marked_failed_not_ready(tmp_path, monkeypatch):
    """Before: 0 chars + no exception => parse_status 'ready' + empty resume."""
    from app.services import parsing

    class _Resume:
        storage_key = "u/scan.pdf"
        raw_text = ""
        structured_data = None
        parse_status = "pending"
        parse_note = ""
        id = None

    class _Repo:
        class db:
            @staticmethod
            def commit():
                pass

        def get_owned(self, *_a, **_k):
            return _Resume()

    monkeypatch.setattr(parsing.storage, "get_object", lambda *_a, **_k: _image_only_pdf())
    svc = parsing.ParsingService.__new__(parsing.ParsingService)
    svc.resumes = _Repo()
    svc.vectors = None

    out = svc.parse_resume("rid", "uid")
    assert out.parse_status == "failed"
    assert "scanned image" in out.parse_note


# ------------------------------------------------------- §1.9 status guard
@pytest.mark.parametrize("status", ["pending", "failed"])
def test_analysis_refuses_an_unparsed_resume(status):
    from app.services.analysis_service import _require_parsed

    class _R:
        parse_status = status
        parse_note = "boom"

    with pytest.raises(ResumeNotReady):
        _require_parsed(_R())


def test_analysis_allows_a_ready_resume():
    from app.services.analysis_service import _require_parsed

    class _R:
        parse_status = "ready"
        parse_note = ""

    _require_parsed(_R())  # must not raise


def test_failed_resume_message_includes_the_parse_note():
    from app.services.analysis_service import _require_parsed

    class _R:
        parse_status = "failed"
        parse_note = "No readable text found - this file may be a scanned image."

    with pytest.raises(ResumeNotReady, match="scanned image"):
        _require_parsed(_R())


# ------------------------------------ the guard must not flunk short text files
@pytest.mark.parametrize(
    "name,text,should_fail",
    [
        ("cv.txt", "Jane Doe\njane@x.com\nEXPERIENCE\nAcme - Eng\n- Did work", False),
        ("cv.txt", "Priya Sharma\npriya@x.com\nExperienced engineer.", False),
        ("cv.txt", "", True),
        ("cv.docx", "Jane Doe jane@x.com Engineer", False),
        ("scan.pdf", "Jane Doe", True),      # a PDF this empty is a scan
        ("real.pdf", " ".join(["word"] * 40), False),
    ],
)
def test_extraction_failure_is_format_aware(name, text, should_fail):
    """"May be a scanned image" is only a sane diagnosis for a PDF."""
    from app.services.parsing import _extraction_failure

    assert bool(_extraction_failure(name, text)) is should_fail


def test_scanned_pdf_note_differs_from_empty_file_note():
    from app.services.parsing import _extraction_failure

    assert "scanned image" in _extraction_failure("scan.pdf", "Jane Doe")
    assert "scanned image" not in _extraction_failure("cv.txt", "")
