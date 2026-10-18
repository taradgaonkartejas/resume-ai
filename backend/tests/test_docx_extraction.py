"""DOCX must not lose content the candidate supplied.

The PDF path got column detection, annotation-layer links and an image-only
check. The DOCX path was `document.paragraphs` plus tables, and five kinds
of content were being dropped. Each was reproduced before it was fixed:

  D1  the Word header/footer band -- "8208300420" was simply absent
  D2  list markers -- python-docx returns `p.text` with the bullet glyph
      stripped, because Word stores it as numbering metadata, so every
      bullet reached `structure_text` looking like a heading
  D3  text boxes -- present in document.xml, invisible to `paragraphs`
  D4  hyperlinks -- `extract_links` early-returned for anything non-PDF, so
      a resume hyperlinking the word "LinkedIn" lost the URL and was then
      docked for `contact.links.missing`
  D5  `parse_facts` recorded only `tables`

Fixtures are BUILT here with python-docx (plus one XML splice for the text
box, which python-docx cannot author) rather than committed as binaries --
the same approach the PDF tests take with reportlab.
"""

from __future__ import annotations

import io
import zipfile

import docx
import pytest

from app.services.parsing import (
    collect_parse_facts,
    extract_links,
    extract_text,
    parse_advisories,
)

HYPERLINK_RT = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)

TEXTBOX_XML = (
    '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    ' xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict>'
    '<v:shape style="width:200pt;height:60pt"><v:textbox><w:txbxContent>'
    "<w:p><w:r><w:t>SIDEBAR Pune, India</w:t></w:r></w:p>"
    "</w:txbxContent></v:textbox></v:shape></w:pict></w:r></w:p>"
)


def _splice_textbox(content: bytes) -> bytes:
    """python-docx cannot author a text box, so one is spliced in directly."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(content)) as zin, zipfile.ZipFile(out, "w") as zo:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = (
                    data.decode("utf8")
                    .replace("</w:body>", TEXTBOX_XML + "</w:body>")
                    .encode("utf8")
                )
            zo.writestr(item, data)
    return out.getvalue()


def _build(header=True, bullets=True, table=True, link=True, textbox=True) -> bytes:
    d = docx.Document()
    if header:
        d.sections[0].header.paragraphs[0].text = (
            "Tejas T | tejas@example.com | +91 8208300420"
        )
    d.add_heading("Work Experience", level=1)
    d.add_paragraph("Software Engineer, Ceinsys Tech Ltd.")
    if bullets:
        d.add_paragraph("Cut p95 response time by 40%", style="List Bullet")
        d.add_paragraph("Lifted throughput to 10 jobs/hour", style="List Bullet")
    if table:
        d.add_heading("Skills", level=1)
        t = d.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "Languages"
        t.rows[0].cells[1].text = "Java, Python"
    if link:
        from docx.oxml.shared import OxmlElement, qn

        para = d.add_paragraph()
        rid = d.part.relate_to(
            "https://www.linkedin.com/in/tejas-t", HYPERLINK_RT, is_external=True
        )
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        node = OxmlElement("w:t")
        node.text = "LinkedIn"
        run.append(node)
        hl.append(run)
        para._p.append(hl)
    buf = io.BytesIO()
    d.save(buf)
    content = buf.getvalue()
    return _splice_textbox(content) if textbox else content


# ------------------------------------------------------------------- D1
class TestHeaderBand:
    def test_header_text_reaches_the_parser(self):
        text = extract_text("cv.docx", _build())
        assert "8208300420" in text

    def test_the_header_comes_first(self):
        """Contact details belong at the top for `structure_text`, and that
        is where they are printed."""
        text = extract_text("cv.docx", _build())
        assert text.splitlines()[0].startswith("Tejas T")

    def test_a_document_without_a_header_is_unaffected(self):
        text = extract_text("cv.docx", _build(header=False))
        assert text.splitlines()[0] == "Work Experience"


# ------------------------------------------------------------------- D2
class TestListMarkers:
    def test_list_paragraphs_regain_their_marker(self):
        text = extract_text("cv.docx", _build())
        assert "- Cut p95 response time by 40%" in text
        assert "- Lifted throughput to 10 jobs/hour" in text

    def test_ordinary_paragraphs_do_not_gain_one(self):
        """The marker is evidence. Adding it everywhere would destroy the
        distinction it exists to carry."""
        text = extract_text("cv.docx", _build())
        assert "- Software Engineer, Ceinsys Tech Ltd." not in text
        assert "- Work Experience" not in text

    def test_bullets_survive_structuring(self):
        """The point of D2: the experience section parses as bullets rather
        than as a pile of headings."""
        from app.services.parsing import structure_text

        data = structure_text(extract_text("cv.docx", _build()))
        bullets = [b for e in data.get("experience", []) for b in e.get("bullets", [])]
        assert any("p95" in b for b in bullets)


# ------------------------------------------------------------------- D3
class TestTextBoxes:
    def test_text_box_content_is_read(self):
        assert "SIDEBAR Pune, India" in extract_text("cv.docx", _build())

    def test_absent_text_boxes_change_nothing(self):
        assert "SIDEBAR" not in extract_text("cv.docx", _build(textbox=False))


# ------------------------------------------------------------------- D4
class TestHyperlinks:
    def test_a_docx_hyperlink_is_extracted(self):
        links = extract_links("cv.docx", _build())
        assert "https://www.linkedin.com/in/tejas-t" in links["links"]

    def test_an_unknown_host_is_not_returned(self):
        """The allowlist is shared with the PDF path. Templates hyperlink
        every project title, and those belong to the project."""
        d = docx.Document()
        d.add_paragraph("x")
        d.part.relate_to("https://example.com/whatever", HYPERLINK_RT, is_external=True)
        buf = io.BytesIO()
        d.save(buf)
        assert extract_links("cv.docx", buf.getvalue())["links"] == []

    def test_a_non_docx_non_pdf_is_still_empty(self):
        assert extract_links("cv.txt", b"hello")["links"] == []

    def test_a_corrupt_docx_does_not_raise(self):
        """Links are a bonus, never a parse failure."""
        assert extract_links("cv.docx", b"not a zip") == {
            "links": [], "email": "", "phone": ""
        }


# ------------------------------------------------------------------- D5
class TestParseFacts:
    def test_facts_cover_the_new_layers(self):
        content = _build()
        facts = collect_parse_facts("cv.docx", content, extract_text("cv.docx", content))
        assert facts["format"] == "docx"
        assert facts["tables"] == 1
        assert facts["text_boxes"] == 1
        assert facts["header_footer_text"] is True

    def test_advisories_warn_about_the_skipped_layers(self):
        content = _build()
        facts = collect_parse_facts("cv.docx", content, extract_text("cv.docx", content))
        joined = " ".join(parse_advisories(facts)).lower()
        assert "header" in joined
        assert "text box" in joined

    def test_a_plain_docx_produces_no_layer_advisories(self):
        """An advisory shown on every upload is an advisory nobody reads."""
        content = _build(header=False, table=False, textbox=False, link=False)
        facts = collect_parse_facts("cv.docx", content, extract_text("cv.docx", content))
        joined = " ".join(parse_advisories(facts)).lower()
        assert "header" not in joined
        assert "text box" not in joined

    def test_collect_never_raises_on_rubbish(self):
        assert collect_parse_facts("cv.docx", b"not a zip", "") is not None


# ------------------------------------------------------- no PDF regression
class TestPdfIsUntouched:
    @pytest.mark.parametrize(
        "name", ["Tejas_Taradgaonkar_Resume_Java.pdf", "Tejas_Taradgaonkar_Resume_react_java.pdf"]
    )
    def test_the_real_pdfs_still_parse_and_warn_about_nothing(self, name):
        import pathlib

        pdf = pathlib.Path("/home/user/uploads") / name
        if not pdf.exists():
            pytest.skip("real resumes not present")
        content = pdf.read_bytes()
        text = extract_text(name, content)
        assert len(text) > 1000
        assert parse_advisories(collect_parse_facts(name, content, text)) == []


# ---------------------------------------------------------------------------
# Regression: the same profile arriving by two routes rendered twice.
#
# Found on the first REAL .docx (samples/Tejas_Resume.docx). The text layer
# printed "github.com/taradgaonkartejas" and D4's relationship reader returned
# "https://github.com/taradgaonkartejas"; merge_extracted_links built its seen
# set from NORMALISED values and probed it with a RAW one, so the two never
# matched and contact.links held four entries for two profiles.
#
# PDFs could not expose this: their text layer shows anchor words, not URLs,
# so there was never anything to collide with.
# ---------------------------------------------------------------------------

from app.services.parsing import _link_identity, merge_extracted_links  # noqa: E402


class TestLinkDeduplication:
    def test_bare_and_qualified_forms_are_one_profile(self):
        data = {"contact": {"links": ["github.com/taradgaonkartejas"]}}
        merge_extracted_links(
            data, {"links": ["https://github.com/taradgaonkartejas"], "email": "", "phone": ""}
        )
        assert data["contact"]["links"] == ["https://github.com/taradgaonkartejas"]

    def test_the_qualified_url_wins_the_value(self):
        """fpdf2 gets this string as a link target; no scheme means no link."""
        data = {"contact": {"links": ["linkedin.com/in/x"]}}
        merge_extracted_links(
            data, {"links": ["https://linkedin.com/in/x"], "email": "", "phone": ""}
        )
        assert data["contact"]["links"][0].startswith("https://")

    def test_the_text_layer_still_wins_on_order(self):
        data = {"contact": {"links": ["github.com/a", "linkedin.com/in/b"]}}
        merge_extracted_links(
            data,
            {"links": ["https://linkedin.com/in/b", "https://github.com/a"],
             "email": "", "phone": ""},
        )
        assert [_link_identity(u) for u in data["contact"]["links"]] == [
            "github.com/a", "linkedin.com/in/b"
        ]

    def test_a_genuinely_new_link_is_still_added(self):
        data = {"contact": {"links": ["github.com/a"]}}
        merge_extracted_links(
            data, {"links": ["https://linkedin.com/in/b"], "email": "", "phone": ""}
        )
        assert len(data["contact"]["links"]) == 2

    def test_www_and_trailing_slash_do_not_create_a_duplicate(self):
        data = {"contact": {"links": ["https://www.linkedin.com/in/x/"]}}
        merge_extracted_links(
            data, {"links": ["https://linkedin.com/in/x"], "email": "", "phone": ""}
        )
        assert len(data["contact"]["links"]) == 1

    def test_two_different_profiles_on_one_host_are_kept(self):
        """Dedup must key on the whole path, not the host."""
        data = {"contact": {"links": ["github.com/alice"]}}
        merge_extracted_links(
            data, {"links": ["https://github.com/bob"], "email": "", "phone": ""}
        )
        assert len(data["contact"]["links"]) == 2
